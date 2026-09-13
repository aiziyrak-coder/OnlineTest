"""O'qituvchilar (faculty) Excel import — JSHSHIR=login=parol, kafedra, lavozim.

Manba: O'qituvchilar ro'yhati.xlsx, sheet `Xodimlar`.

  python manage.py import_faculty --file "data/oqituvchilar.xlsx" --limit 10
  python manage.py import_faculty --file "data/oqituvchilar.xlsx" --apply
  python manage.py import_faculty --file "data/oqituvchilar.xlsx" --apply --photos
"""
from __future__ import annotations

import re

import bcrypt
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from apps.core.models import AppUser, Kafedra

HEADER_ALIASES = {
    "pinfl": ["jshshir-kod", "jshshir kod", "jshshir", "pinfl"],
    "kafedra": ["c", "kafedra", "кафедра"],
    "first_name": ["ismi", "ism", "имя"],
    "last_name": ["familiya", "фамилия"],
    "middle_name": ["otasining ismi", "otasi", "отчество"],
    "stavka": ["stavka", "ставка"],
    "position": ["lavozim", "должность"],
}


def normalize_header(v) -> str:
    return str(v or "").strip().lower()


def cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _hash_pw(plain: str) -> str:
    return bcrypt.hashpw(plain.encode("utf-8"), bcrypt.gensalt(rounds=10)).decode("utf-8")


def _full_name(last_name: str, first_name: str, middle_name: str) -> str:
    parts = [p for p in (last_name, first_name, middle_name) if p]
    return " ".join(parts).strip() or first_name or last_name


def _sheet_xml_name(wb, sheet_name: str) -> str:
    """openpyxl sheet → xl/worksheets/sheetN.xml (rich-data uchun)."""
    try:
        from openpyxl.workbook.workbook import Workbook  # noqa: F401

        idx = wb.sheetnames.index(sheet_name) + 1
        return f"xl/worksheets/sheet{idx}.xml"
    except Exception:
        return "xl/worksheets/sheet1.xml"


class Command(BaseCommand):
    help = "O'qituvchilarni Excel'dan import qiladi (role=faculty, JSHSHIR=login=parol)."

    def add_arguments(self, parser):
        parser.add_argument("--file", required=True, help="Excel fayl yo'li (.xlsx)")
        parser.add_argument("--sheet", default="Xodimlar", help="Sheet nomi (default: Xodimlar)")
        parser.add_argument("--limit", type=int, default=0, help="Nechta qator (0 = hammasi)")
        parser.add_argument("--apply", action="store_true", help="Haqiqatan saqlash (default: dry-run)")
        parser.add_argument(
            "--photos",
            action="store_true",
            help="Rich-data (Picture in Cell) rasmlarni ham o'rnatish",
        )
        parser.add_argument(
            "--reset-password",
            action="store_true",
            help="Mavjud foydalanuvchi parolini ham JSHSHIR ga qayta o'rnatish",
        )

    def handle(self, *args, **opts):
        try:
            import openpyxl
        except ImportError:
            raise CommandError("openpyxl o'rnatilmagan: pip install openpyxl")

        file_path = opts["file"]
        sheet_name = str(opts["sheet"] or "Xodimlar").strip()
        apply_changes = bool(opts["apply"])
        with_photos = bool(opts["photos"])
        reset_password = bool(opts["reset_password"])
        limit = int(opts["limit"])

        wb = openpyxl.load_workbook(file_path, data_only=True)
        if sheet_name not in wb.sheetnames:
            raise CommandError(f"Sheet topilmadi: {sheet_name}. Mavjud: {wb.sheetnames}")
        ws = wb[sheet_name]

        header_row = [normalize_header(c.value) for c in ws[1]]
        col_index: dict[str, int] = {}
        for field, aliases in HEADER_ALIASES.items():
            for i, h in enumerate(header_row):
                if h in aliases:
                    col_index[field] = i
                    break
        if "pinfl" not in col_index:
            raise CommandError(f"JSHSHIR ustuni topilmadi. Sarlavhalar: {header_row}")
        if "kafedra" not in col_index:
            raise CommandError(f"Kafedra ustuni topilmadi. Sarlavhalar: {header_row}")

        def get(row_vals: list, field: str) -> str:
            i = col_index.get(field)
            return cell_str(row_vals[i]) if i is not None and i < len(row_vals) else ""

        rich = None
        if with_photos:
            from apps.core.management.commands.seed_student_photos import (
                RichImageIndex,
                to_jpeg_data_uri,
            )

            sheet_xml = _sheet_xml_name(wb, sheet_name)
            rich = RichImageIndex(file_path, sheet_xml)
            if not rich.available:
                self.stdout.write(self.style.WARNING("Rich-data topilmadi — rasmlar o'tkazib yuboriladi."))
                rich = None
            else:
                self.stdout.write(f"Rasm ustuni: {rich.image_col} (sheet xml={sheet_xml})")

        created = 0
        updated = 0
        skipped = 0
        photos_set = 0
        kafedra_cache: dict[str, Kafedra] = {}

        mode = "APPLY" if apply_changes else "DRY-RUN"
        self.stdout.write(f"[{mode}] sheet={sheet_name} rows≈{ws.max_row - 1}")

        with transaction.atomic():
            for r in range(2, ws.max_row + 1):
                if limit and (created + updated + skipped) >= limit:
                    break
                row_vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
                pinfl = re.sub(r"\D", "", get(row_vals, "pinfl"))
                if not pinfl or len(pinfl) < 10:
                    skipped += 1
                    continue
                kafedra_name = get(row_vals, "kafedra")
                if not kafedra_name:
                    skipped += 1
                    continue
                name = _full_name(
                    get(row_vals, "last_name"),
                    get(row_vals, "first_name"),
                    get(row_vals, "middle_name"),
                )
                if not name:
                    skipped += 1
                    continue
                position = get(row_vals, "position")[:200]
                stavka = get(row_vals, "stavka")[:64]

                kafedra = kafedra_cache.get(kafedra_name.lower())
                if kafedra is None:
                    kafedra, _ = Kafedra.objects.get_or_create(
                        name=kafedra_name,
                        defaults={"is_active": True},
                    )
                    kafedra_cache[kafedra_name.lower()] = kafedra

                profile_b64 = ""
                if rich is not None:
                    result = rich.image_bytes_for_row(r)
                    if result is not None:
                        raw, _ext = result
                        uri = to_jpeg_data_uri(raw)
                        if uri:
                            profile_b64 = uri

                user = AppUser.objects.filter(pk=pinfl).first()
                if user is None:
                    if apply_changes:
                        AppUser.objects.create(
                            id=pinfl,
                            password=_hash_pw(pinfl),
                            role="faculty",
                            name=name,
                            status="Active",
                            kafedra=kafedra,
                            position=position,
                            stavka=stavka,
                            profile_image=profile_b64,
                        )
                    created += 1
                    if profile_b64:
                        photos_set += 1
                else:
                    if apply_changes:
                        user.role = "faculty"
                        user.name = name
                        user.kafedra = kafedra
                        user.position = position
                        user.stavka = stavka
                        if reset_password:
                            user.password = _hash_pw(pinfl)
                        if profile_b64 and (
                            not user.profile_image or len(user.profile_image) < 50
                        ):
                            user.profile_image = profile_b64
                            photos_set += 1
                        user.save()
                    updated += 1

            if not apply_changes:
                transaction.set_rollback(True)

        self.stdout.write(
            self.style.SUCCESS(
                f"created={created} updated={updated} skipped={skipped} photos={photos_set}"
            )
        )
        if not apply_changes:
            self.stdout.write("Dry-run — DB o'zgarmadi. Yozish uchun --apply qo'shing.")
