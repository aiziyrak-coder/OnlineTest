"""HEMIS eksportidan pasport raqamlarini hisoblarga bog'laydi.

Test markazida kirish pasport raqami bilan bo'ladi (12 xonali talaba ID va
alohida parol o'rniga). Pasport bizning bazada saqlanmasdi — shu buyruq uni
HEMIS Excel faylidan olib qo'yadi.

Moslashtirish tartibi:
  1) "Talaba ID" ustuni bo'lsa — eng ishonchlisi, shu bo'yicha;
  2) bo'lmasa — ism-familiya bo'yicha (imlo farqlari hisobga olinadi);
  3) ikki kishi bir xil ismli bo'lsa — qo'lda hal qilish uchun ro'yxatga chiqadi.

    python manage.py import_passports --file hemis.xlsx            # sinov
    python manage.py import_passports --file hemis.xlsx --apply    # yozish
    python manage.py import_passports --file hemis.xlsx --role ordinator --apply
"""
from __future__ import annotations

import re
import unicodedata

from django.core.management.base import BaseCommand, CommandError

from apps.core.models import AppUser

HEADER = {
    "passport": ["pasport raqami", "pasport", "passport", "паспорт"],
    "student_id": ["talaba id", "talaba_id", "student id", "id"],
    "full_name": ["to‘liq ismi", "to'liq ismi", "toliq ismi", "f.i.sh", "fish", "ismi", "фио"],
}
PASSPORT_RE = re.compile(r"^[A-Z]{2}\d{7}$")


def norm_header(v) -> str:
    return str(v or "").strip().lower()


def norm_name(v) -> str:
    """Ism solishtirish uchun: apostrof, x/h va qo'shimchalar farqi yo'qoladi."""
    s = unicodedata.normalize("NFKD", str(v or "")).lower()
    for a in ("‘", "’", "`", "ʻ", "ʼ", "'"):
        s = s.replace(a, "")
    s = s.replace("x", "h").replace("ʼ", "")
    s = re.sub(r"[^a-zа-яё ]+", " ", s)
    skip = {"ogli", "oglu", "qizi", "kizi", "ugli", "o g li"}
    return " ".join(p for p in s.split() if p not in skip and len(p) > 1)


class Command(BaseCommand):
    help = "HEMIS Excel faylidan pasport raqamlarini hisoblarga bog'laydi."

    def add_arguments(self, parser):
        parser.add_argument("--file", required=True)
        parser.add_argument("--role", default="", help="Faqat shu toifadagi hisoblar")
        parser.add_argument("--apply", action="store_true", help="Haqiqatan yozish")

    def handle(self, *args, **o):
        try:
            from openpyxl import load_workbook
        except ImportError as ex:  # pragma: no cover
            raise CommandError("openpyxl kerak: pip install openpyxl") from ex

        wb = load_workbook(o["file"], read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = [norm_header(c) for c in next(rows)]
        col: dict[str, int] = {}
        for field, names in HEADER.items():
            for i, h in enumerate(header):
                if h in names:
                    col[field] = i
                    break
        if "passport" not in col:
            raise CommandError("Faylda 'Pasport raqami' ustuni topilmadi. Ustunlar: %s"
                               % ", ".join(h for h in header if h))
        if "student_id" not in col and "full_name" not in col:
            raise CommandError("Faylda 'Talaba ID' yoki ism ustuni bo'lishi kerak.")

        users = AppUser.objects.all()
        if o["role"]:
            users = users.filter(role=o["role"])
        by_id = {str(u.pk): u for u in users}
        by_name: dict[str, list] = {}
        for u in users:
            by_name.setdefault(norm_name(u.name), []).append(u)

        matched, skipped, ambiguous, bad = [], [], [], []
        seen_passport: dict[str, str] = {}
        for r in rows:
            if not r:
                continue
            passport = re.sub(r"[^A-Za-z0-9]", "", str(r[col["passport"]] or "")).upper()
            sid = str(r[col["student_id"]] or "").strip() if "student_id" in col else ""
            name = str(r[col["full_name"]] or "").strip() if "full_name" in col else ""
            if not PASSPORT_RE.match(passport):
                if passport:
                    bad.append("%s (%s)" % (passport, name or sid))
                continue
            u = by_id.get(sid)
            if u is None and name:
                found = by_name.get(norm_name(name)) or []
                if len(found) == 1:
                    u = found[0]
                elif len(found) > 1:
                    ambiguous.append(name)
                    continue
            if u is None:
                skipped.append(name or sid or passport)
                continue
            if passport in seen_passport and seen_passport[passport] != str(u.pk):
                ambiguous.append("%s — bitta pasport ikki kishida" % passport)
                continue
            seen_passport[passport] = str(u.pk)
            matched.append((u, passport))

        self.stdout.write("Faylda mos keldi: %d | topilmadi: %d | ikki xil: %d | noto'g'ri raqam: %d"
                          % (len(matched), len(skipped), len(ambiguous), len(bad)))
        for lbl, lst in (("topilmadi", skipped), ("ikki xil", ambiguous), ("noto'g'ri", bad)):
            for x in lst[:10]:
                self.stdout.write("   %s: %s" % (lbl, x))

        changed = [(u, p) for u, p in matched if (u.passport or "") != p]
        self.stdout.write("Yoziladi: %d ta (qolganida allaqachon bor)" % len(changed))
        for u, p in changed[:5]:
            self.stdout.write("   %s -> %s (%s)" % (u.pk, p, u.name[:34]))
        if not o["apply"]:
            self.stdout.write("\nKo'rish rejimi. Yozish uchun --apply qo'shing.")
            return
        for u, p in changed:
            u.passport = p
        AppUser.objects.bulk_update([u for u, _ in changed], ["passport"], batch_size=500)
        self.stdout.write("YOZILDI: %d ta hisobga pasport bog'landi." % len(changed))
